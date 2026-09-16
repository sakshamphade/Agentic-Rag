import io

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


class ChatDownloader:

    @staticmethod
    def build_chat(messages):

        chat = ""

        for message in messages:

            role = message["role"].capitalize()

            chat += f"{role}:\n"

            chat += message["content"]

            chat += "\n\n"

        return chat

    # ----------------------------
    # TXT
    # ----------------------------

    @staticmethod
    def create_txt(messages):

        chat = ChatDownloader.build_chat(messages)

        return chat.encode("utf-8")

    # ----------------------------
    # DOCX
    # ----------------------------

    @staticmethod
    def create_docx(messages):

        document = Document()

        document.add_heading(
            "Agentic AI Chat",
            level=1
        )

        for message in messages:

            role = message["role"].capitalize()

            document.add_heading(
                role,
                level=2
            )

            document.add_paragraph(
                message["content"]
            )

        buffer = io.BytesIO()

        document.save(buffer)

        buffer.seek(0)

        return buffer

    # ----------------------------
    # PDF
    # ----------------------------

    @staticmethod
    def create_pdf(messages):

        buffer = io.BytesIO()

        pdf = canvas.Canvas(
            buffer,
            pagesize=letter
        )

        width, height = letter

        y = height - 40

        pdf.setFont(
            "Helvetica",
            11
        )

        pdf.drawString(
            40,
            y,
            "Agentic AI Chat"
        )

        y -= 30

        for message in messages:

            role = message["role"].capitalize()

            pdf.drawString(
                40,
                y,
                role + ":"
            )

            y -= 18

            text = message["content"]

            words = text.split()

            line = ""

            for word in words:

                if len(line + word) < 90:

                    line += word + " "

                else:

                    pdf.drawString(
                        60,
                        y,
                        line
                    )

                    y -= 18

                    line = word + " "

                if y < 40:

                    pdf.showPage()

                    pdf.setFont(
                        "Helvetica",
                        11
                    )

                    y = height - 40

            pdf.drawString(
                60,
                y,
                line
            )

            y -= 30

        pdf.save()

        buffer.seek(0)

        return buffer